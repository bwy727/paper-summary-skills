#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
answering.py — 文献问答工具

子命令：
  dump [--root OUT_DIR]
      检测 keyword_summary.xlsx / paper_archiving.xlsx 是否存在，并输出其内容 JSON：
      {"has_keyword":bool, "has_archiving":bool,
       "keyword": {"papers":[{"文献名","keywords":[...]}]},
       "archiving": {"columns":[...], "rows":[{...}]}}
      （供 Claude 据问题语义筛选文献）

  scanmd [--folder IN_DIR]
      无表时的兜底：递归读取 md 全文，输出 [{"filename","stem","text"},...]

  genmd --out 关键词.md --data names.json
      names.json = {"question":str, "direct":[{"name","note"}...], "related":[...]}
      生成文献名单 markdown。

  filter --archiving paper_archiving.xlsx --names names.json --out 关键词_文献总结筛选.xlsx
      从文献总结表中筛出这些文献名对应的行，另存为新表（保留表头与格式）。

默认路径：输入 IN_DIR = D:/AIacademic/ReadPaper/notes_done，
输出 OUT_DIR = D:/AIacademic/ReadPaper/xlsx_summary。
只给文件名（不含目录）的 --out / --archiving 一律落在 OUT_DIR 内。
"""
import argparse
import importlib
import json
import subprocess
import sys
from pathlib import Path

# Windows 控制台默认 GBK，文献中的非 GBK 字符会导致 print 崩溃 —— 强制 UTF-8。
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass


def _ensure(pip_name, mod_name):
    try:
        importlib.import_module(mod_name)
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", pip_name, "-q"])


_ensure("openpyxl", "openpyxl")

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

# summary_note 生成的阅读笔记统一命名为 `[note] <论文标题>.md`，该前缀不计入文献名。
NOTE_PREFIX = "[note]"
# 递归时跳过的工作目录与生成物
SKIP_DIRS = {"_work", ".git"}
SKIP_FILES = {"searching_readme.md"}
# 输入/输出路径约定
IN_DIR = r"D:\AIacademic\ReadPaper\notes_done"
OUT_DIR = r"D:\AIacademic\ReadPaper\xlsx_summary"



def resolve_out(p, default_dir=OUT_DIR):
    """只给文件名（不含目录）时落在输出文件夹内；给了路径则按原样使用。"""
    q = Path(str(p))
    if q.is_absolute() or q.parent != Path("."):
        return q
    return Path(default_dir) / q


def norm(s):
    return str(s).strip()


def stem_key(name):
    """文献名键：md 文件名去扩展名，并去掉开头的 `[note] ` 前缀。"""
    stem = norm(Path(str(name)).stem)
    if stem.lower().startswith(NOTE_PREFIX):
        stem = stem[len(NOTE_PREFIX):].strip()
    return stem


def md_to_text(path):
    return Path(str(path)).read_text(encoding="utf-8", errors="replace")


def list_md(folder):
    root = Path(folder)
    if not root.exists():
        return []
    out = []
    for p in root.rglob("*.md"):
        if p.name.startswith(".") or p.name in SKIP_FILES:
            continue
        if any(part in SKIP_DIRS for part in p.relative_to(root).parts[:-1]):
            continue
        out.append(p)
    return sorted(out)


def read_keyword(xlsx):
    wb = load_workbook(str(xlsx))
    ws = wb.active
    header = [c.value for c in ws[1]]
    kws = [h for h in header[1:] if h]
    papers = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        hit = [h for h, v in zip(header[1:], r[1:]) if h and v]
        papers.append({"文献名": r[0], "keywords": hit})
    wb.close()
    return {"keywords": kws, "papers": papers}


def read_archiving(xlsx):
    wb = load_workbook(str(xlsx))
    ws = wb.active
    header = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        rows.append({h: v for h, v in zip(header, r)})
    wb.close()
    return {"columns": header, "rows": rows}


def _emit(obj, json_out, summary):
    if json_out:
        Path(json_out).write_text(json.dumps(obj, ensure_ascii=False), encoding="utf-8")
        summary = dict(summary)
        summary["json_out"] = json_out
        print(json.dumps(summary, ensure_ascii=False))
    else:
        print(json.dumps(obj, ensure_ascii=False))


def _find_table(root, name):
    """先在 root 下找，再在 root/xlsx_summary 下找（兼容传 --root . 的旧调用）。"""
    for p in (root / name, root / "xlsx_summary" / name):
        if p.exists():
            return p
    return root / name


def cmd_dump(args):
    root = Path(args.root)
    kw_path = _find_table(root, "keyword_summary.xlsx")
    ar_path = _find_table(root, "paper_archiving.xlsx")
    out = {"has_keyword": kw_path.exists(), "has_archiving": ar_path.exists()}
    if kw_path.exists():
        out["keyword"] = read_keyword(kw_path)
    if ar_path.exists():
        out["archiving"] = read_archiving(ar_path)
    _emit(out, args.json_out,
          {"has_keyword": out["has_keyword"], "has_archiving": out["has_archiving"]})


def cmd_scanmd(args):
    res = [{"filename": p.name, "stem": stem_key(p.name), "text": md_to_text(p)}
           for p in list_md(args.folder)]
    _emit(res, args.json_out, {"count": len(res)})


def item_name(it):
    """条目可为字符串（仅文件名）或 {"name","note"}。"""
    return it.get("name", "") if isinstance(it, dict) else str(it)


def item_note(it):
    return it.get("note", "") if isinstance(it, dict) else ""


def _md_section(lines, heading, items, note_label):
    lines.append("## %s" % heading)
    lines.append("")
    if not items:
        lines.append("（无）")
    else:
        for it in items:
            lines.append("- **%s**" % item_name(it))
            note = item_note(it)
            if note:
                lines.append("  - %s%s" % (note_label, note))
    lines.append("")


def cmd_genmd(args):
    data = json.loads(Path(args.data).read_text(encoding="utf-8"))
    direct = data.get("direct", [])
    related = data.get("related", [])
    out = resolve_out(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    lines = ["# %s" % data.get("question", "文献名单"), ""]
    _md_section(lines, "可直接回答的文献", direct, "如何回答：")
    _md_section(lines, "话题相关的文献", related, "相关之处：")
    out.write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps({"out": str(out), "direct": len(direct), "related": len(related)},
                     ensure_ascii=False))


def cmd_filter(args):
    names = json.loads(Path(args.names).read_text(encoding="utf-8"))
    if isinstance(names, dict):
        items = names.get("direct", []) + names.get("related", [])
        wanted = set(stem_key(item_name(it)) for it in items)
    else:
        wanted = set(stem_key(item_name(it)) for it in names)
    src = load_workbook(str(resolve_out(args.archiving)))
    sws = src.active
    header = [c.value for c in sws[1]]
    wb = Workbook()
    ws = wb.active
    ws.title = "筛选结果"
    ws.append(header)
    kept = 0
    for r in sws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        if stem_key(r[0]) in wanted:
            ws.append(list(r))
            kept += 1
    # 复用总结表的排版
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    for ci in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 30
        hc = ws.cell(row=1, column=ci)
        hc.font = Font(bold=True)
        hc.fill = header_fill
        hc.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for rr in range(2, ws.max_row + 1):
        for cc in range(1, len(header) + 1):
            ws.cell(row=rr, column=cc).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    src.close()
    out = resolve_out(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(json.dumps({"out": str(out), "kept": kept}, ensure_ascii=False))


def main():
    ap = argparse.ArgumentParser(description="文献问答工具")
    sub = ap.add_subparsers(dest="cmd", required=True)

    d = sub.add_parser("dump")
    d.add_argument("--root", default=OUT_DIR)
    d.add_argument("--json-out", dest="json_out", default=None)
    d.set_defaults(func=cmd_dump)

    s = sub.add_parser("scanmd")
    s.add_argument("--folder", default=IN_DIR)
    s.add_argument("--json-out", dest="json_out", default=None)
    s.set_defaults(func=cmd_scanmd)

    g = sub.add_parser("genmd")
    g.add_argument("--out", default="文献名单.md")
    g.add_argument("--data", required=True)
    g.set_defaults(func=cmd_genmd)

    f = sub.add_parser("filter")
    f.add_argument("--archiving", default="paper_archiving.xlsx")
    f.add_argument("--names", required=True)
    f.add_argument("--out", required=True)
    f.set_defaults(func=cmd_filter)

    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
