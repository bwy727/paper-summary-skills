#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
archiving.py — 文献总结表（paper_archiving.xlsx）工具

子命令：
  extract --folder notes_done --out xlsx_summary/paper_archiving.xlsx --json-out work.json
      递归列出 folder 下尚未收录的 .md，把结果写入 --json-out（UTF-8）：
      {"existing": [...已收录文献名...], "new": [{"filename","stem","text"}, ...]}
      并在 stdout 打印简短摘要（供 Claude 阅读 work.json 后逐篇生成总结）。
      不传 --json-out 则直接打印完整 JSON 到 stdout。
      文献名取 md 文件名（去扩展名），并去掉开头的 `[note] ` 前缀
      （summary_note 生成的阅读笔记统一命名为 `[note] <论文标题>.md`）。

  write --out paper_archiving.xlsx --data rows.json
      rows.json 为 [{"文献名","发表年份","期刊","摘要","研究内容","主要结果","研究思路"}, ...]
      追加写入（按文献名去重，忽略重复），单元格自动换行、表头加粗、冻结首行。

年份与期刊只依据文献本身（md 文件名与笔记正文）判断，本脚本不读取原 PDF。
"""
import argparse
import importlib
import json
import subprocess
import sys
from pathlib import Path

# Windows 控制台默认 GBK，文献中的非 GBK 字符会导致 print 崩溃 —— 强制 UTF-8。
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass


def _ensure(pip_name, mod_name):
    try:
        importlib.import_module(mod_name)
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", pip_name, "-q"])


_ensure("openpyxl", "openpyxl")

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

COLUMNS = ["文献名", "发表年份", "期刊", "摘要", "研究内容", "主要结果", "研究思路"]

# summary_note 生成的阅读笔记统一命名为 `[note] <论文标题>.md`，该前缀不计入文献名。
NOTE_PREFIX = "[note]"
# 递归时跳过的工作目录与生成物
SKIP_DIRS = {"_work", ".git"}
SKIP_FILES = {"searching_readme.md"}
# 输入/输出路径约定
IN_DIR = r"D:\AIacademic\ReadPaper\notes_done"
OUT_DIR = r"D:\AIacademic\ReadPaper\xlsx_summary"



def norm(s):
    return str(s).strip()


def stem_key(name):
    """文献名键：md 文件名去扩展名，并去掉开头的 `[note] ` 前缀。"""
    stem = norm(Path(str(name)).stem)
    if stem.lower().startswith(NOTE_PREFIX):
        stem = stem[len(NOTE_PREFIX):].strip()
    return stem


def md_to_text(path):
    return Path(str(path)).read_text(encoding="utf-8", errors="replace")


def list_md(folder):
    root = Path(folder)
    if not root.exists():
        return []
    out = []
    for p in root.rglob("*.md"):
        if p.name.startswith(".") or p.name in SKIP_FILES:
            continue
        if any(part in SKIP_DIRS for part in p.relative_to(root).parts[:-1]):
            continue
        out.append(p)
    return sorted(out)


def existing_names(xlsx):
    p = Path(xlsx)
    if not p.exists():
        return set()
    wb = load_workbook(str(p), read_only=True)
    ws = wb.active
    names = set()
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if row and row[0]:
            names.add(stem_key(row[0]))
    wb.close()
    return names


def emit(obj, json_out, summary):
    """大体量结果写入文件并打印摘要；否则整段打印。"""
    if json_out:
        Path(json_out).write_text(json.dumps(obj, ensure_ascii=False), encoding="utf-8")
        summary = dict(summary)
        summary["json_out"] = json_out
        print(json.dumps(summary, ensure_ascii=False))
    else:
        print(json.dumps(obj, ensure_ascii=False))


def cmd_extract(args):
    have = existing_names(args.out)
    new = []
    for p in list_md(args.folder):
        if stem_key(p.name) in have:
            continue
        new.append({"filename": p.name, "stem": stem_key(p.name), "text": md_to_text(p)})
    emit({"existing": sorted(have), "new": new}, args.json_out,
         {"existing_count": len(have), "new_count": len(new)})


def _format(ws):
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    widths = {"文献名": 32, "发表年份": 10, "期刊": 18, "摘要": 40,
              "研究内容": 36, "主要结果": 44, "研究思路": 44}
    for ci, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = widths.get(col, 24)
        hc = ws.cell(row=1, column=ci)
        hc.font = Font(bold=True)
        hc.fill = header_fill
        hc.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def cmd_write(args):
    data = json.loads(Path(args.data).read_text(encoding="utf-8"))
    p = Path(args.out)
    if p.exists():
        wb = load_workbook(str(p))
        ws = wb.active
        have = existing_names(args.out)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "文献总结"
        ws.append(COLUMNS)
        have = set()
    added = 0
    for row in data:
        key = stem_key(row.get("文献名", ""))
        if not key or key in have:
            continue
        ws.append([row.get(c, "") for c in COLUMNS])
        have.add(key)
        added += 1
    _format(ws)
    wb.save(str(p))
    print(json.dumps({"added": added, "total_rows": ws.max_row - 1, "out": str(p)},
                     ensure_ascii=False))


def main():
    ap = argparse.ArgumentParser(description="文献总结表工具")
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("extract")
    e.add_argument("--folder", default=IN_DIR)
    e.add_argument("--out", default=str(Path(OUT_DIR) / "paper_archiving.xlsx"))
    e.add_argument("--json-out", dest="json_out", default=None)
    e.set_defaults(func=cmd_extract)
    w = sub.add_parser("write")
    w.add_argument("--out", default=str(Path(OUT_DIR) / "paper_archiving.xlsx"))
    w.add_argument("--data", required=True)
    w.set_defaults(func=cmd_write)
    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
