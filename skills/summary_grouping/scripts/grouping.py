#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
grouping.py — 文献分组表（keyword_summary.xlsx，关键词矩阵）工具

子命令：
  extract --folder note --out keyword_summary.xlsx
      输出 JSON：
      {"existing_keywords":[...], "existing_papers":[...], "new":[{"filename","stem","text"},...]}
      （供 Claude 复用已有关键词、为新文献提关键词，保持关键词一致性）

  write --out keyword_summary.xlsx --data kw.json
      kw.json 为 [{"文献名":"xx.docx","keywords":["a","b",...]}, ...]
      建立/更新矩阵：第一列文献名，其后每个关键词一列，文献含该关键词则标 X。
      开启自动筛选、冻结首列与表头、自动换行；并在同目录生成 searching_readme.md。
"""
import argparse
import importlib
import json
import subprocess
import sys
from pathlib import Path

# Windows 控制台默认 GBK，docx 中的非 GBK 字符会导致 print 崩溃 —— 强制 UTF-8。
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass


def _ensure(pip_name, mod_name):
    try:
        importlib.import_module(mod_name)
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", pip_name, "-q"])


_ensure("python-docx", "docx")
_ensure("openpyxl", "openpyxl")

import docx  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402


def norm(s):
    return str(s).strip()


def stem_key(name):
    return norm(Path(str(name)).stem)


def docx_to_text(path):
    d = docx.Document(str(path))
    parts = []
    for p in d.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def list_docx(folder):
    root = Path(folder)
    if not root.exists():
        return []
    return sorted(p for p in root.rglob("*.docx") if not p.name.startswith("~$"))


def load_matrix(xlsx):
    """返回 (keywords:list, rows:list[dict])；dict 含 '文献名' 与每个命中关键词=X。"""
    p = Path(xlsx)
    if not p.exists():
        return [], []
    wb = load_workbook(str(p))
    ws = wb.active
    header = [c.value for c in ws[1]] if ws.max_row >= 1 else ["文献名"]
    keywords = [h for h in header[1:] if h]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        rec = {"文献名": r[0]}
        for h, v in zip(header[1:], r[1:]):
            if h and v:
                rec[h] = "X"
        rows.append(rec)
    wb.close()
    return keywords, rows


def cmd_extract(args):
    keywords, rows = load_matrix(args.out)
    existing_papers = [stem_key(r["文献名"]) for r in rows]
    new = []
    for p in list_docx(args.folder):
        if stem_key(p.name) in existing_papers:
            continue
        new.append({"filename": p.name, "stem": p.stem, "text": docx_to_text(p)})
    obj = {
        "existing_keywords": keywords,
        "existing_papers": [r["文献名"] for r in rows],
        "new": new,
    }
    if args.json_out:
        Path(args.json_out).write_text(json.dumps(obj, ensure_ascii=False), encoding="utf-8")
        print(json.dumps({"existing_keywords_count": len(keywords),
                          "new_count": len(new), "json_out": args.json_out},
                         ensure_ascii=False))
    else:
        print(json.dumps(obj, ensure_ascii=False))


def _format(ws, header):
    fill = PatternFill("solid", fgColor="E2EFDA")
    ws.column_dimensions["A"].width = 36
    for ci in range(2, len(header) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 14
    for ci, h in enumerate(header, start=1):
        c = ws.cell(row=1, column=ci)
        c.font = Font(bold=True)
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        for ci in range(2, len(header) + 1):
            ws.cell(row=r, column=ci).alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(header)), ws.max_row)


README_TEXT = """# keyword_summary 检索使用说明

`keyword_summary.xlsx` 是一张关键词矩阵：第一列为文献名，其后每一列对应一个关键词，
某文献含某关键词则该单元格标记为 `X`。借助 Excel 自动筛选即可按关键词检索文献。

## 一、按单个关键词检索
1. 用 Excel 打开 `keyword_summary.xlsx`。
2. 第一行表头已开启自动筛选（每个表头单元格右侧有下拉箭头）。
   若未显示，选中第一行后点击「数据」→「筛选」。
3. 找到目标关键词所在的列，点击该列表头的下拉箭头。
4. 在筛选面板中取消勾选「（空白）」，仅保留 `X`，点击「确定」。
5. 此时表格只剩下含该关键词的文献行，第一列即为对应文献名。

## 二、按多个关键词组合检索（同时满足，AND）
1. 先按「一、」对第一个关键词列筛选出 `X`。
2. 不要清除上一步筛选，继续对第二个关键词列重复同样操作（仅保留 `X`）。
3. 依次对更多关键词列操作，最终显示的文献会同时含所有已筛选的关键词。

## 三、清除筛选 / 恢复全部
- 清除单列筛选：点击该列下拉箭头 →「从…中清除筛选」。
- 一次清除全部筛选：点击「数据」→「清除」。

## 四、提示
- 关键词在不同文献间保持同一表述，因此同一关键词只对应一列，便于检索。
- 首列与表头已冻结，左右/上下滚动时文献名与关键词始终可见。
- 新增文献后重新运行「整理文献标签」会增量更新本表，已有内容不受影响。
"""


def cmd_write(args):
    data = json.loads(Path(args.data).read_text(encoding="utf-8"))
    keywords, rows = load_matrix(args.out)
    by_stem = {stem_key(r["文献名"]): r for r in rows}
    for item in data:
        fn = str(item.get("文献名", ""))
        if not fn:
            continue
        key = stem_key(fn)
        rec = by_stem.get(key)
        if rec is None:
            rec = {"文献名": fn}
            by_stem[key] = rec
            rows.append(rec)
        for kw in item.get("keywords", []):
            kw = norm(kw)
            if not kw:
                continue
            if kw not in keywords:
                keywords.append(kw)
            rec[kw] = "X"

    wb = Workbook()
    ws = wb.active
    ws.title = "文献分组"
    header = ["文献名"] + keywords
    ws.append(header)
    for rec in rows:
        ws.append([rec.get("文献名", "")] + ["X" if rec.get(k) == "X" else "" for k in keywords])
    _format(ws, header)
    wb.save(args.out)
    (Path(args.out).resolve().parent / "searching_readme.md").write_text(
        README_TEXT, encoding="utf-8")
    print(json.dumps({"papers": len(rows), "keywords": len(keywords), "out": args.out,
                      "readme": "searching_readme.md"}, ensure_ascii=False))


def main():
    ap = argparse.ArgumentParser(description="文献分组表工具")
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("extract")
    e.add_argument("--folder", default="note")
    e.add_argument("--out", default="keyword_summary.xlsx")
    e.add_argument("--json-out", dest="json_out", default=None)
    e.set_defaults(func=cmd_extract)
    w = sub.add_parser("write")
    w.add_argument("--out", default="keyword_summary.xlsx")
    w.add_argument("--data", required=True)
    w.set_defaults(func=cmd_write)
    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
